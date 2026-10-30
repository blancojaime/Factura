"""Exportación a Excel (.xlsx) con openpyxl: importes numéricos, formato profesional y configuración de impresión."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .engine import ReporteDef, Tabla, Texto

AZUL = "1F3A5F"
CLARO = "DEE7F3"
ZEBRA = "F6F8FB"
FMT = {"money": "#,##0.00;[Red]-#,##0.00", "num": "#,##0.00", "int": "#,##0", "pct": '0.00"%"'}


def exportar(rep: ReporteDef, ruta: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = (rep.titulo[:28] or "Reporte").replace("/", "-").replace("\\", "-").replace(":", "-")
    ncols = max([len(t.cols) for t in rep.tablas] or [4])
    r = 1
    ws.cell(r, 1, rep.empresa).font = Font(bold=True, size=13, color=AZUL)
    ws.cell(r, ncols, rep.gestion).font = Font(bold=True, color=AZUL)
    ws.cell(r, ncols).alignment = Alignment(horizontal="right")
    r += 1
    if rep.nit:
        ws.cell(r, 1, f"NIT: {rep.nit}").font = Font(color="6C757D", size=9)
        r += 1
    if rep.direccion:
        ws.cell(r, 1, rep.direccion).font = Font(color="6C757D", size=9)
        r += 1
    r += 1
    ws.cell(r, 1, rep.titulo.upper()).font = Font(bold=True, size=15, color=AZUL)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    ws.cell(r, 1).alignment = Alignment(horizontal="center")
    r += 1
    for linea in (rep.subtitulo, rep.periodo, rep.moneda):
        if linea:
            ws.cell(r, 1, linea).font = Font(italic=True, color="6C757D")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
            ws.cell(r, 1).alignment = Alignment(horizontal="center")
            r += 1
    r += 1
    delgado = Side(style="thin", color="D5DBE4")
    primera_cabecera = None
    anchos = [10.0] * ncols
    for b in rep.bloques:
        if isinstance(b, Texto) and b.estilo == "salto":
            r += 2
            continue
        if isinstance(b, Texto):
            c = ws.cell(r, 1, b.texto)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(bold=b.estilo == "titulo", color=AZUL if b.estilo == "titulo" else "000000", size=11 if b.estilo == "titulo" else 10)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
            ws.row_dimensions[r].height = max(15, 15 * (len(b.texto) // 110 + 1))
            r += 1
            continue
        t: Tabla = b
        if t.titulo:
            ws.cell(r, 1, t.titulo).font = Font(bold=True, color=AZUL, size=11)
            r += 1
        if not t.sin_encabezado:
            for j, c in enumerate(t.cols, 1):
                x = ws.cell(r, j, c.titulo)
                x.font = Font(bold=True, color="FFFFFF")
                x.fill = PatternFill("solid", fgColor=AZUL)
                x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                anchos[j - 1] = max(anchos[j - 1], min(len(c.titulo) + 2, 30))
            if primera_cabecera is None:
                primera_cabecera = r
            ws.row_dimensions[r].height = 30
            r += 1
        for i, f in enumerate(t.filas):
            for j, c in enumerate(t.cols, 1):
                v = f.valores[j - 1] if j - 1 < len(f.valores) else None
                x = ws.cell(r, j, v if v != "" else None)
                if c.tipo in FMT:
                    x.number_format = FMT[c.tipo] if c.tipo != "num" else "#,##0." + "0" * c.decimales
                    x.alignment = Alignment(horizontal="right")
                elif j == 1 and f.nivel:
                    x.alignment = Alignment(indent=f.nivel)
                if f.estilo in ("titulo", "total", "subtotal", "grupo"):
                    x.font = Font(bold=True)
                if f.estilo in ("titulo", "total"):
                    x.fill = PatternFill("solid", fgColor=CLARO)
                elif i % 2 == 1 and f.estilo == "":
                    x.fill = PatternFill("solid", fgColor=ZEBRA)
                if f.estilo == "total":
                    x.border = Border(top=Side(style="medium", color=AZUL))
                else:
                    x.border = Border(bottom=delgado)
                texto = str(v) if v is not None else ""
                anchos[j - 1] = max(anchos[j - 1], min(len(texto) + 2 + (f.nivel * 2 if j == 1 else 0), 60))
            r += 1
        r += 1
    for n in rep.notas:
        ws.cell(r, 1, n).font = Font(italic=True, color="6C757D", size=9)
        r += 1
    for j, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = a
    if primera_cabecera:
        ws.freeze_panes = ws.cell(primera_cabecera + 1, 1)
    ws.page_setup.orientation = "landscape" if rep.orientacion == "H" else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{primera_cabecera}:{primera_cabecera}" if primera_cabecera else None
    ws.oddFooter.center.text = "Página &P de &N"
    ws.oddFooter.left.text = "ContaBol"
    wb.properties.title = rep.titulo
    wb.properties.creator = "ContaBol"
    wb.save(ruta)
    return ruta
