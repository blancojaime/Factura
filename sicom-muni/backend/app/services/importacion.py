"""Importacion desde Excel (.xlsx): presupuesto por apertura programatica, Catalogo CHB y objetivos del POA.

Todas las importaciones tienen dos fases: (1) validar y mostrar un resumen sin tocar la base (`aplicar=False`);
(2) aplicar (`aplicar=True`), que solo procede si el archivo no tiene errores (todo o nada).
"""
import io
import re
import unicodedata
from collections.abc import Callable
from datetime import date
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..enums import EstadoPreventivo
from ..errors import ReglaNegocioError
from ..models import CatalogoCHB, ObjetivoPOA, PartidaPresupuestaria, PreventivoC31
from . import poa

MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 20000
CENT = Decimal("0.01")


# ------------------------------------------------------------------ utilidades de lectura
def norm(texto) -> str:
    t = unicodedata.normalize("NFD", str(texto if texto is not None else "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", t.lower()).strip()


def a_decimal(valor) -> Decimal | None:
    """Convierte celdas numericas o textos como 'Bs 1.234,56' / '1,234.56' a Decimal (2 decimales)."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, bool):
        raise ValueError("valor no numérico")
    if isinstance(valor, int | float | Decimal):
        return Decimal(str(valor)).quantize(CENT, ROUND_HALF_UP)
    s = re.sub(r"(?i)bs\.?|\s| ", "", str(valor))
    if not s:
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        mil = "." if dec == "," else ","
        s = s.replace(mil, "").replace(dec, ".")
    elif "," in s:
        s = s.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    elif "." in s and re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s).quantize(CENT, ROUND_HALF_UP)
    except InvalidOperation as exc:
        raise ValueError("valor no numérico") from exc


def a_codigo(valor) -> str:
    """Codigo como texto: 30111505.0 -> '30111505'; ' 20 ' -> '20'."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


Regla = tuple[str, bool, Callable[[str], bool]]  # (campo, requerido, predicado sobre el encabezado normalizado)


def _mapear(fila, reglas: list[Regla]) -> dict[str, int] | None:
    mapa: dict[str, int] = {}
    usadas: set[int] = set()
    encabezados = [norm(c) for c in fila]
    for campo, _req, pred in reglas:
        for i, h in enumerate(encabezados):
            if i not in usadas and h and pred(h):
                mapa[campo] = i
                usadas.add(i)
                break
    if all(c in mapa for c, req, _ in reglas if req):
        return mapa
    return None


def _leer(contenido: bytes, nombre: str, reglas: list[Regla], ayuda: str) -> tuple[list[tuple], dict[str, int], int]:
    """Devuelve (filas de datos, mapa de columnas, N° de fila del primer dato). Busca el encabezado en todas las hojas."""
    if not contenido:
        raise ReglaNegocioError("El archivo está vacío", "ARCHIVO_INVALIDO")
    if len(contenido) > MAX_BYTES:
        raise ReglaNegocioError("El archivo supera los 5 MB", "ARCHIVO_INVALIDO")
    if not nombre.lower().endswith((".xlsx", ".xlsm")):
        raise ReglaNegocioError("Suba un archivo Excel con extensión .xlsx (si tiene .xls, ábralo en Excel y use «Guardar como» .xlsx)",
                                "ARCHIVO_INVALIDO")
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - archivo corrupto o no es un xlsx
        raise ReglaNegocioError("No se pudo leer el archivo Excel; verifique que no esté dañado ni protegido con contraseña",
                                "ARCHIVO_INVALIDO") from exc
    try:
        for hoja in libro.worksheets:
            filas = list(hoja.iter_rows(values_only=True, max_row=MAX_FILAS + 40))
            for idx, fila in enumerate(filas[:40]):
                mapa = _mapear(fila, reglas)
                if mapa is not None:
                    return filas[idx + 1:], mapa, idx + 2
    finally:
        libro.close()
    raise ReglaNegocioError("No se encontró la fila de encabezados. " + ayuda, "ENCABEZADOS_NO_ENCONTRADOS")


def _celda(fila: tuple, mapa: dict[str, int], campo: str):
    i = mapa.get(campo)
    return fila[i] if i is not None and i < len(fila) else None


def _vacia(fila: tuple) -> bool:
    return all(c is None or (isinstance(c, str) and not c.strip()) for c in fila)


def _resultado(tipo: str, aplicar: bool, total: int, errores: list[dict], **kw) -> dict:
    if aplicar and errores:
        raise ReglaNegocioError("El archivo tiene errores: corríjalos y vuelva a importar (no se guardó nada)",
                                "IMPORTACION_CON_ERRORES", [f"Fila {e['fila']}: {e['mensaje']}" for e in errores[:50]])
    return {"tipo": tipo, "aplicado": aplicar, "total_filas": total, "errores": errores, **kw}


# ------------------------------------------------------------------ 1. Presupuesto por apertura programatica
def _r(campo, req, pred) -> Regla:
    return (campo, req, pred)


REGLAS_PRESUPUESTO: list[Regla] = [
    _r("descripcion_partida", False, lambda h: "descripcion" in h and "partida" in h),
    _r("descripcion_apertura", True, lambda h: ("descripcion" in h and "apertura" in h) or h in ("descripcion", "denominacion")),
    _r("apertura", True, lambda h: ("apertura" in h and "descripcion" not in h) or h in ("codigo", "cod")),
    _r("partida", True, lambda h: h.startswith("partida") or h in ("cod partida", "codigo partida")),
    _r("fuente", True, lambda h: h.startswith("fuente") or h in ("ff", "fte")),
    _r("organismo", True, lambda h: h.startswith("organismo") or h in ("of", "org")),
    _r("presupuesto", True, lambda h: h.startswith("presupuesto")),
    _r("ejecutado", True, lambda h: h.startswith("ejecut")),
    _r("saldo", True, lambda h: h.startswith("saldo")),
]
AYUDA_PRESUPUESTO = ("Columnas requeridas: código (apertura programática), descripción apertura, partida, fuente, organismo, "
                     "presupuesto, ejecutado y saldo. Use la plantilla de importación.")


def parsear_apertura(texto: str) -> tuple[str, str, str]:
    partes = [p for p in re.split(r"[^0-9A-Za-z]+", texto.strip()) if p]
    if len(partes) != 3:
        raise ValueError(f"código de apertura «{texto}» inválido: use Programa-Proyecto-Actividad separados por guion (ej. 002-0000-001)")
    return partes[0], partes[1], partes[2]


def importar_presupuesto(db: Session, contenido: bytes, nombre: str, gestion: int, aplicar: bool) -> dict:
    filas, mapa, primera = _leer(contenido, nombre, REGLAS_PRESUPUESTO, AYUDA_PRESUPUESTO)
    errores: list[dict] = []
    advertencias: list[str] = []
    validas: dict[tuple, dict] = {}
    omitidas = 0
    ult_apertura = ult_desc = ""
    total = 0
    for n, fila in enumerate(filas, start=primera):
        if _vacia(fila):
            continue
        apertura_txt = a_codigo(_celda(fila, mapa, "apertura"))
        if norm(apertura_txt).startswith("total"):
            continue
        total += 1
        try:
            if apertura_txt:
                ult_apertura, ult_desc = apertura_txt, str(_celda(fila, mapa, "descripcion_apertura") or "").strip() or ult_desc
            elif not ult_apertura:
                raise ValueError("falta el código de apertura programática")
            programa, proyecto, actividad = parsear_apertura(ult_apertura)
            partida = a_codigo(_celda(fila, mapa, "partida"))
            if not re.fullmatch(r"\d{5}", partida):
                raise ValueError(f"partida «{partida}» inválida (5 dígitos)")
            if not "10000" <= partida <= "49999":
                omitidas += 1  # el sistema solo usa 1xxxx a 4xxxx
                continue
            fuente, organismo = a_codigo(_celda(fila, mapa, "fuente")), a_codigo(_celda(fila, mapa, "organismo"))
            if not fuente or not organismo:
                raise ValueError("faltan fuente u organismo")
            monto = a_decimal(_celda(fila, mapa, "presupuesto"))
            ejec = a_decimal(_celda(fila, mapa, "ejecutado"))
            saldo = a_decimal(_celda(fila, mapa, "saldo"))
            if monto is None or ejec is None or saldo is None:
                raise ValueError("presupuesto, ejecutado y saldo son obligatorios")
            if min(monto, ejec, saldo) < 0:
                raise ValueError("los montos no pueden ser negativos")
            if abs(monto - ejec - saldo) > CENT:
                advertencias.append(f"Fila {n}: presupuesto − ejecutado ({monto - ejec}) no coincide con el saldo ({saldo}); se usa el saldo del archivo")
            clave = (partida, programa, proyecto, actividad, fuente, organismo)
            if clave in validas:
                raise ValueError("combinación apertura/partida/fuente/organismo repetida en el archivo")
            validas[clave] = {
                "fila": n, "apertura": ult_apertura, "descripcion_apertura": ult_desc[:250],
                "descripcion_partida": str(_celda(fila, mapa, "descripcion_partida") or "").strip()[:200],
                "monto": monto, "ejecutado": ejec, "saldo": saldo}
        except ValueError as exc:
            errores.append({"fila": n, "mensaje": str(exc)})

    existentes = {(p.codigo_partida, p.programa, p.proyecto, p.actividad, p.fuente, p.organismo): p
                  for p in db.scalars(select(PartidaPresupuestaria).where(PartidaPresupuestaria.gestion == gestion))}
    nuevas = actualizadas = 0
    for clave, v in validas.items():
        existente = existentes.get(clave)
        if existente is None:
            nuevas += 1
        else:
            actualizadas += 1
            activos = db.scalar(select(func.coalesce(func.sum(PreventivoC31.importe), 0)).where(
                PreventivoC31.partida_id == existente.id, PreventivoC31.estado == EstadoPreventivo.ACTIVO.value))
            if activos and Decimal(activos) > 0:
                advertencias.append(f"Partida {clave[0]} ({'-'.join(clave[1:4])}): el sistema tiene preventivos activos por Bs {Decimal(activos):,.2f}; "
                                    "su saldo se reemplaza por el del Excel (verifique que ya los incluya)")
    res = _resultado("presupuesto", aplicar, total, errores, nuevas=nuevas, actualizadas=actualizadas, omitidas=omitidas,
                     advertencias=advertencias[:100],
                     muestra=[{"apertura": v["apertura"], "partida": k[0], "fuente": k[4], "organismo": k[5],
                               "presupuesto": str(v["monto"]), "ejecutado": str(v["ejecutado"]), "saldo": str(v["saldo"])}
                              for k, v in list(validas.items())[:8]])
    if aplicar:
        for clave, v in validas.items():
            p = existentes.get(clave)
            if p is None:
                p = PartidaPresupuestaria(
                    codigo_partida=clave[0], programa=clave[1], proyecto=clave[2], actividad=clave[3], fuente=clave[4],
                    organismo=clave[5], gestion=gestion, descripcion=v["descripcion_partida"] or f"Partida {clave[0]}")
                db.add(p)
            elif v["descripcion_partida"]:
                p.descripcion = v["descripcion_partida"]
            p.descripcion_apertura = v["descripcion_apertura"]
            p.monto_aprobado, p.monto_ejecutado, p.saldo_disponible = v["monto"], v["ejecutado"], v["saldo"]
        db.flush()
    return res


# ------------------------------------------------------------------ 2. Catalogo CHB
REGLAS_CHB: list[Regla] = [
    _r("descripcion", True, lambda h: any(k in h for k in ("descripcion", "denominacion", "detalle", "producto"))),
    _r("fecha", False, lambda h: "fecha" in h),
    _r("unidad", False, lambda h: "unidad" in h),
    _r("precio", True, lambda h: "precio" in h),
    _r("codigo", True, lambda h: any(k in h for k in ("unspsc", "unesco", "codigo", "cod"))),
]
AYUDA_CHB = "Columnas requeridas: código UNSPSC, descripción del código y precio (opcionales: unidad y fecha). Use la plantilla."


def importar_catalogo_chb(db: Session, contenido: bytes, nombre: str, aplicar: bool) -> dict:
    filas, mapa, primera = _leer(contenido, nombre, REGLAS_CHB, AYUDA_CHB)
    errores: list[dict] = []
    advertencias: list[str] = []
    validas: dict[str, dict] = {}
    total = 0
    for n, fila in enumerate(filas, start=primera):
        if _vacia(fila):
            continue
        total += 1
        try:
            codigo = a_codigo(_celda(fila, mapa, "codigo"))
            if not re.fullmatch(r"\d{8}", codigo):
                raise ValueError(f"código UNSPSC «{codigo}» inválido (8 dígitos)")
            desc = str(_celda(fila, mapa, "descripcion") or "").strip()
            if len(desc) < 3:
                raise ValueError("falta la descripción del código")
            precio = a_decimal(_celda(fila, mapa, "precio"))
            if precio is not None and precio < 0:
                raise ValueError("el precio no puede ser negativo")
            if precio is None:
                advertencias.append(f"Fila {n}: el código {codigo} no tiene precio")
            if codigo in validas:
                advertencias.append(f"Fila {n}: el código {codigo} está repetido en el archivo; prevalece la última fila")
            validas[codigo] = {"descripcion": desc[:250], "precio": precio, "unidad": str(_celda(fila, mapa, "unidad") or "").strip()[:40]}
        except ValueError as exc:
            errores.append({"fila": n, "mensaje": str(exc)})
    existentes = {c.codigo_unspsc: c for c in db.scalars(select(CatalogoCHB).where(CatalogoCHB.codigo_unspsc.in_(list(validas))))} if validas else {}
    nuevos = len([c for c in validas if c not in existentes])
    reemplazados = len(validas) - nuevos
    res = _resultado("catalogo_chb", aplicar, total, errores, nuevas=nuevos, actualizadas=reemplazados, omitidas=0,
                     advertencias=advertencias[:100], fecha_registro=date.today().isoformat(),
                     muestra=[{"codigo": c, "descripcion": v["descripcion"], "precio": str(v["precio"]) if v["precio"] is not None else "",
                               "accion": "REEMPLAZA" if c in existentes else "NUEVO"} for c, v in list(validas.items())[:8]])
    if aplicar:
        hoy = date.today()
        for codigo, v in validas.items():
            c = existentes.get(codigo)
            if c is None:
                c = CatalogoCHB(codigo_unspsc=codigo)
                db.add(c)
            c.descripcion_bien, c.precio_referencial_nacional, c.fecha_registro, c.activo = v["descripcion"], v["precio"], hoy, True
            if v["unidad"]:
                c.unidad_medida = v["unidad"]
        db.flush()
    return res


# ------------------------------------------------------------------ 3. Objetivos del POA por programa
REGLAS_POA: list[Regla] = [
    _r("denominacion", False, lambda h: h.startswith("denominaci")),
    _r("codigo_objetivo", True, lambda h: ("cod" in h and "obj" in h) or h in ("codigo objetivo", "cod obj")),
    _r("descripcion_objetivo", True, lambda h: h.startswith("descripc") and "obj" in h),
    _r("programa", True, lambda h: h in ("prog", "programa", "cod prog", "codigo programa", "cod programa")),
]
AYUDA_POA = "Columnas requeridas: PROG, COD OBJ y DESCRIPCIÓN OBJETIVO (opcional: DENOMINACIÓN DEL PROGRAMA). Use la plantilla."


def importar_objetivos_poa(db: Session, contenido: bytes, nombre: str, gestion: int, aplicar: bool) -> dict:
    filas, mapa, primera = _leer(contenido, nombre, REGLAS_POA, AYUDA_POA)
    errores: list[dict] = []
    validas: dict[str, dict] = {}
    total = 0
    for n, fila in enumerate(filas, start=primera):
        if _vacia(fila):
            continue
        total += 1
        try:
            programa = a_codigo(_celda(fila, mapa, "programa"))
            if not programa or len(programa) > 10:
                raise ValueError("programa vacío o demasiado largo")
            cod = str(_celda(fila, mapa, "codigo_objetivo") or "").strip()
            desc = str(_celda(fila, mapa, "descripcion_objetivo") or "").strip()
            if not cod or not desc:
                raise ValueError("faltan el código o la descripción del objetivo")
            clave = poa.normalizar_programa(programa)
            if clave in validas:
                raise ValueError(f"el programa {programa} está repetido en el archivo")
            validas[clave] = {"programa": programa, "denominacion": str(_celda(fila, mapa, "denominacion") or "").strip()[:250],
                              "codigo": cod[:30], "descripcion": desc}
        except ValueError as exc:
            errores.append({"fila": n, "mensaje": str(exc)})
    existentes = {poa.normalizar_programa(o.programa): o for o in db.scalars(select(ObjetivoPOA).where(ObjetivoPOA.gestion == gestion))}
    nuevos = len([k for k in validas if k not in existentes])
    res = _resultado("objetivos_poa", aplicar, total, errores, nuevas=nuevos, actualizadas=len(validas) - nuevos, omitidas=0,
                     advertencias=[], muestra=[{"programa": v["programa"], "objetivo": v["codigo"], "descripcion": v["descripcion"][:80]}
                                               for v in list(validas.values())[:8]])
    if aplicar:
        for clave, v in validas.items():
            o = existentes.get(clave)
            if o is None:
                o = ObjetivoPOA(gestion=gestion, programa=v["programa"])
                db.add(o)
            o.denominacion_programa, o.codigo_objetivo, o.descripcion_objetivo = v["denominacion"], v["codigo"], v["descripcion"]
        db.flush()
    return res


# ------------------------------------------------------------------ plantillas descargables
def plantilla(tipo: str) -> bytes:
    libro = Workbook()
    hoja = libro.active
    if tipo == "presupuesto":
        hoja.title = "Presupuesto"
        hoja.append(["Código apertura programática", "Descripción apertura", "Partida", "Descripción partida (opcional)",
                     "Fuente", "Organismo", "Presupuesto", "Ejecutado", "Saldo"])
        hoja.append(["002-0000-001", "ADMINISTRACIÓN CENTRAL - Gestión administrativa", "34200", "Productos minerales y no metálicos", "20", "230", 80000, 12000, 68000])
        hoja.append(["002-0000-001", "ADMINISTRACIÓN CENTRAL - Gestión administrativa", "25800", "Mantenimiento y reparación", "20", "230", 45000, 0, 45000])
    elif tipo == "chb":
        hoja.title = "CHB"
        hoja.append(["Código UNSPSC", "Descripción código", "Precio", "Unidad (opcional)"])
        hoja.append(["30111505", "Cemento portland IP-30 (bolsa de 50 kg)", 58.00, "Bolsa"])
        hoja.append(["14111507", "Papel bond tamaño carta 75 g", 32.00, "Resma"])
    elif tipo == "poa":
        hoja.title = "POA"
        hoja.append(["PROG", "DENOMINACIÓN DEL PROGRAMA", "COD OBJ", "DESCRIPCIÓN OBJETIVO"])
        hoja.append(["002", "ADMINISTRACIÓN CENTRAL", "OBJ3", "Garantizar la gestión administrativa eficiente de la entidad"])
    else:
        raise ReglaNegocioError("Plantilla desconocida", "PLANTILLA_NO_EXISTE")
    for col in hoja.columns:
        hoja.column_dimensions[col[0].column_letter].width = 34
    buf = io.BytesIO()
    libro.save(buf)
    return buf.getvalue()
