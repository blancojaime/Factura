"""Informe Excel de la validación / importación."""
from __future__ import annotations

import datetime as dt
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .importer import ERROR, WARN, Importer, Result

HFILL = PatternFill("solid", fgColor="1F3864")
HFONT = Font(bold=True, color="FFFFFF")
RED = PatternFill("solid", fgColor="F8CBAD")
YEL = PatternFill("solid", fgColor="FFF2CC")

CALC_COLS = [("fila", "Fila"), ("unidad", "Unidad"), ("codigo", "Código"), ("grupo", "Grupo contable"),
             ("estado", "Estado"), ("incorporacion", "Incorporación"), ("vida_util", "Vida útil"),
             ("costo_historico", "Costo histórico"), ("factor", "Factor UFV"),
             ("costo_actualizado_inicial", "Costo actualizado inicial"),
             ("actualizacion_gestion", "Actualización gestión"), ("costo_total_actualizado", "Costo total actualizado"),
             ("dep_gestion", "Dep. gestión"), ("dep_acumulada_inicial", "Dep. acum. inicial"),
             ("actualizacion_dep_acumulada", "Actualiz. dep. acum."), ("dep_acumulada", "Dep. acumulada"),
             ("valor_neto", "Valor neto"), ("dias", "Días")]
MONEY = {"costo_historico", "costo_actualizado_inicial", "actualizacion_gestion", "costo_total_actualizado",
         "dep_gestion", "dep_acumulada_inicial", "actualizacion_dep_acumulada", "dep_acumulada", "valor_neto",
         "costo_anterior", "valor_actualizado", "dep_acumulada_anterior", "valor_neto_anterior", "nuevo_costo",
         "costo", "costo_actualizado"}


def _header(ws, row: int, titles) -> None:
    for j, t in enumerate(titles, start=1):
        c = ws.cell(row, j, t)
        c.font, c.fill = HFONT, HFILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, minw: int = 9, maxw: int = 60) -> None:
    for col in ws.columns:
        L = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(minw, min(maxw, L + 2))


def write_report(path: str, imp: Importer, res: Result, wb_path: str, mode: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"
    ws["A1"] = f"INFORME DE {mode.upper()} — Importador Excel → VSIAF"
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    info = [
        ("Fecha", dt.datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Archivo Excel", wb_path),
        ("Entidad", f"{imp.state.entidad} — {imp.db.entidad_nombre}"),
        ("Carpeta VSIAF", imp.db.install_dir),
        ("Unidad por defecto", imp.opt.default_unit or "(ninguna)"),
        ("Fecha de cálculo (vista previa)", imp.opt.calc_date.strftime("%d/%m/%Y")),
        ("Fuente de UFV", f"{imp.ufv.source}  (último cambio de índice: "
                          f"{imp.ufv.last_change:%d/%m/%Y})" if imp.ufv.last_change else imp.ufv.source),
        ("Resultado", "IMPORTADO" if res.applied else ("VALIDADO (no se escribió nada)" if mode == "validación" else "NO IMPORTADO")),
        ("Errores", res.errors), ("Avisos", res.warnings),
        ("Respaldo", res.backup_path or "—"),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    r0 = 3 + len(info) + 1
    _header(ws, r0, ["Hoja", "Filas leídas", "A importar", "Con error / omitidas"])
    for i, name in enumerate(imp.rows_total, start=r0 + 1):
        ws.cell(i, 1, name)
        ws.cell(i, 2, imp.rows_total[name])
        ws.cell(i, 3, res.counts.get(name, 0))
        ws.cell(i, 4, res.skipped.get(name, 0))
    r1 = r0 + len(imp.rows_total) + 3
    ws.cell(r1 - 1, 1, "TOTALES POR GRUPO CONTABLE (activos nuevos, vista previa a la fecha de cálculo)").font = Font(bold=True)
    _header(ws, r1, ["Grupo contable", "Cant.", "Costo histórico", "Costo actualizado inicial", "Actualización gestión",
                     "Costo total actualizado", "Dep. gestión", "Actualiz. dep. acum.", "Dep. acumulada", "Valor neto"])
    groups: "OrderedDict[str, list]" = OrderedDict()
    for r in sorted(res.computed.get("ACTIVOS", []), key=lambda x: x["codcont"]):
        g = groups.setdefault(r["grupo"], [0] + [0.0] * 8)
        g[0] += 1
        for j, k in enumerate(["costo_historico", "costo_actualizado_inicial", "actualizacion_gestion",
                               "costo_total_actualizado", "dep_gestion", "actualizacion_dep_acumulada",
                               "dep_acumulada", "valor_neto"], start=1):
            g[j] += r[k]
    row = r1 + 1
    tot = [0] + [0.0] * 8
    for name, vals in groups.items():
        ws.cell(row, 1, name)
        for j, v in enumerate(vals, start=2):
            ws.cell(row, j, round(v, 2) if j > 2 else v)
            tot[j - 2] += v
        row += 1
    if groups:
        ws.cell(row, 1, "TOTAL").font = Font(bold=True)
        for j, v in enumerate(tot, start=2):
            c = ws.cell(row, j, round(v, 2) if j > 2 else v)
            c.font = Font(bold=True)
    for r in ws.iter_rows(min_row=r1 + 1, max_row=row, min_col=3, max_col=10):
        for c in r:
            c.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    for L in "CDEFGHIJ":
        ws.column_dimensions[L].width = 18

    wi = wb.create_sheet("INCIDENCIAS")
    _header(wi, 1, ["Hoja", "Fila", "Columna", "Nivel", "Mensaje"])
    for i, it in enumerate(res.issues, start=2):
        for j, v in enumerate([it.sheet, it.row or "", it.column, it.level, it.message], start=1):
            wi.cell(i, j, v)
        fill = RED if it.level == ERROR else (YEL if it.level == WARN else None)
        if fill:
            for j in range(1, 6):
                wi.cell(i, j).fill = fill
    wi.freeze_panes = "A2"
    _autosize(wi)
    wi.column_dimensions["E"].width = 110

    def dump(title: str, rows, cols):
        if not rows:
            return
        w = wb.create_sheet(title)
        _header(w, 1, [c[1] for c in cols])
        for i, r in enumerate(rows, start=2):
            for j, (k, _) in enumerate(cols, start=1):
                c = w.cell(i, j, r.get(k))
                if k in MONEY:
                    c.number_format = "#,##0.00"
                elif isinstance(r.get(k), dt.date):
                    c.number_format = "dd/mm/yyyy"
        w.freeze_panes = "A2"
        _autosize(w)

    dump("CALCULO_ACTIVOS", res.computed.get("ACTIVOS"), CALC_COLS)
    rv = res.computed.get("REVALUOS")
    if rv:
        dump("REVALUOS", rv, [(k, k.replace("_", " ").capitalize()) for k in rv[0]])
    tr = res.computed.get("TRANSFERENCIAS")
    if tr:
        dump("TRANSFERENCIAS", tr, [(k, k.replace("_", " ").capitalize()) for k in tr[0]])
    bj = res.computed.get("BAJAS")
    if bj:
        dump("BAJAS", bj, [(k, k.replace("_", " ").capitalize()) for k in bj[0]])
    wb.save(path)
