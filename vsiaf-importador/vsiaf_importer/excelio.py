"""Generación de la plantilla Excel y lectura del libro de importación."""
from __future__ import annotations

import datetime as dt
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .schema import SHEET_BY_NAME, SHEET_ORDER, SHEETS

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
REQ_FILL = PatternFill("solid", fgColor="C55A11")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")


# ============================================================================= lectura
def norm_header(s: Any) -> str:
    t = unicodedata.normalize("NFD", str(s or "").strip().lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = t.replace("*", "")
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


@dataclass
class Row:
    row: int                    # número de fila en Excel
    data: Dict[str, Any]


@dataclass
class WorkbookData:
    path: str
    sheets: Dict[str, List[Row]] = field(default_factory=dict)
    problems: List[str] = field(default_factory=list)    # problemas de estructura (hojas/columnas)

    def rows(self, sheet: str) -> List[Row]:
        return self.sheets.get(sheet, [])


def _empty(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def read_workbook(path: str) -> WorkbookData:
    wb = load_workbook(path, read_only=True, data_only=True)
    out = WorkbookData(path)
    present = {s.upper(): s for s in wb.sheetnames}
    known = False
    for name in SHEET_ORDER:
        if name not in present:
            continue
        known = True
        ws = wb[present[name]]
        spec = SHEET_BY_NAME[name]
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration:
            out.sheets[name] = []
            continue
        cols = {}
        for i, h in enumerate(header):
            key = norm_header(h)
            if key:
                cols[key] = i
        for c in spec.cols:
            if c.required and c.name not in cols:
                out.problems.append(f"Hoja {name}: falta la columna obligatoria «{c.name}».")
        valid = {c.name for c in spec.cols}
        extra = [k for k in cols if k not in valid]
        if extra:
            out.problems.append(f"Hoja {name}: columnas no reconocidas (se ignoran): {', '.join(extra)}.")
        rows: List[Row] = []
        for n, vals in enumerate(it, start=2):
            data = {k: (vals[i] if i < len(vals) else None) for k, i in cols.items() if k in valid}
            if all(_empty(v) for v in data.values()):
                continue
            rows.append(Row(n, data))
        out.sheets[name] = rows
    wb.close()
    if not known:
        out.problems.append("El libro no contiene ninguna hoja reconocida "
                            f"({', '.join(SHEET_ORDER)}). Use la plantilla del importador.")
    return out


# ----------------------------------------------------------------------------- conversión de celdas
def as_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "SI" if v else "NO"
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def as_number(v: Any) -> Optional[float]:
    if _empty(v):
        return None
    if isinstance(v, bool):
        raise ValueError("no es un número")
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("Bs", "").replace("bs", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        raise ValueError(f"«{v}» no es un número válido") from None


def as_date(v: Any) -> Optional[dt.date]:
    if _empty(v):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 < float(v) < 80000:                         # número de serie de Excel
            return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
        raise ValueError(f"«{v}» no es una fecha válida")
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"«{v}» no es una fecha válida (use dd/mm/aaaa)")


# ============================================================================= plantilla
def _catalog_sheet(wb: Workbook, grupos, estados, motivos, organismos, departamentos) -> Dict[str, str]:
    ws = wb.create_sheet("CATALOGOS")
    ws.sheet_properties.tabColor = "7F7F7F"
    blocks = [
        ("GRUPO CONTABLE", [g["nombre"] for g in grupos],
         [("codigo", [g["codcont"] for g in grupos]), ("vida util (años)", [g["vida_util"] for g in grupos])]),
        ("ESTADO DEL BIEN", [e["nombre"] for e in estados], []),
        ("MOTIVO DE BAJA", [m["nombre"] for m in motivos], [("codigo", [m["codigo"] for m in motivos])]),
        ("ORGANISMO FINANCIADOR", [o["codigo"] for o in organismos],
         [("descripcion", [o["nombre"] for o in organismos])]),
        ("EXPEDIDO", [d["sigla"] for d in departamentos], [("departamento", [d["nombre"] for d in departamentos])]),
        ("ESTADO REGISTRO", ["ELABORADO", "APROBADO"], []),
        ("SI / NO", ["NO", "SI"], []),
    ]
    refs: Dict[str, str] = {}
    col = 1
    for title, values, extras in blocks:
        ws.cell(1, col, title).font = HEADER_FONT
        ws.cell(1, col).fill = HEADER_FILL
        for i, v in enumerate(values, start=2):
            ws.cell(i, col, v)
        letter = get_column_letter(col)
        refs[title] = f"CATALOGOS!${letter}$2:${letter}${len(values) + 1}"
        ws.column_dimensions[letter].width = 38 if title in ("GRUPO CONTABLE", "MOTIVO DE BAJA") else 16
        col += 1
        for (et, ev) in extras:
            ws.cell(1, col, et).font = HEADER_FONT
            ws.cell(1, col).fill = HEADER_FILL
            for i, v in enumerate(ev, start=2):
                ws.cell(i, col, v)
            ws.column_dimensions[get_column_letter(col)].width = 40 if et == "descripcion" else 16
            col += 1
        col += 1
    ws.freeze_panes = "A2"
    return refs


LEEME = [
    ("IMPORTADOR EXCEL → VSIAF", "title"),
    ("Cada hoja de este libro corresponde a una tabla del VSIAF. Llene solo las hojas que necesite; las demás pueden quedar vacías.", ""),
    ("La fila 1 de cada hoja contiene los encabezados (no cambiarlos ni reordenarlos). Los encabezados en naranja son OBLIGATORIOS.", ""),
    ("Pase el cursor sobre un encabezado para ver la ayuda de la columna. La hoja CATALOGOS lista los valores válidos.", ""),
    ("", ""),
    ("ORDEN DE PROCESO", "h"),
    ("1. UNIDADES → 2. OFICINAS → 3. RESPONSABLES → 4. AUXILIARES → 5. ACTIVOS → 6. REVALUOS → 7. TRANSFERENCIAS → 8. BAJAS", ""),
    ("Las hojas se procesan en ese orden, por lo que una fila de ACTIVOS puede usar una oficina/responsable/auxiliar creado en este mismo libro.", ""),
    ("", ""),
    ("REGLAS PRINCIPALES (las mismas que aplica el VSIAF)", "h"),
    ("• El código del activo es único por unidad (se guarda en MAYÚSCULAS, máx. 15 caracteres).", ""),
    ("• La fecha de incorporación no puede ser futura y debe tener índice UFV registrado.", ""),
    ("• La vida útil es la del grupo contable (D.S. 24051). Solo con Incorporación Especial autorizada por la DGCF puede ser distinta.", ""),
    ("• Los activos se importan ELABORADOS (hay que aprobarlos en el VSIAF) salvo que indique estado_registro = APROBADO.", ""),
    ("• Revalúos, transferencias y bajas solo se aplican a activos APROBADOS.", ""),
    ("• Fechas: dd/mm/aaaa. Números: sin separador de miles (10000.50).", ""),
    ("", ""),
    ("DESPUÉS DE IMPORTAR", "h"),
    ("Cierre el importador, abra el VSIAF y ejecute Administradores → Re-indexar. Luego revise el Reporte 6 (Resumen por grupo contable).", ""),
]


def build_template(path: str, grupos, estados, motivos, organismos, departamentos,
                   ejemplos: Optional[Dict[str, List[Dict[str, Any]]]] = None) -> None:
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "LEEME"
    ws0.sheet_properties.tabColor = "1F3864"
    for i, (txt, kind) in enumerate(LEEME, start=1):
        c = ws0.cell(i, 1, txt)
        if kind == "title":
            c.font = Font(bold=True, size=16, color="1F3864")
        elif kind == "h":
            c.font = Font(bold=True, size=12, color="C55A11")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 130
    refs = _catalog_sheet(wb, grupos, estados, motivos, organismos, departamentos)

    dv_map = {
        ("ACTIVOS", "grupo_contable"): refs["GRUPO CONTABLE"],
        ("AUXILIARES", "grupo_contable"): refs["GRUPO CONTABLE"],
        ("ACTIVOS", "estado_bien"): refs["ESTADO DEL BIEN"],
        ("ACTIVOS", "estado_registro"): refs["ESTADO REGISTRO"],
        ("ACTIVOS", "incorporacion_especial"): refs["SI / NO"],
        ("ACTIVOS", "organismo_financiador"): refs["ORGANISMO FINANCIADOR"],
        ("BAJAS", "motivo_baja"): refs["MOTIVO DE BAJA"],
        ("RESPONSABLES", "expedido"): refs["EXPEDIDO"],
    }
    date_cols = {"fecha_incorporacion", "fecha_revaluo", "fecha_baja"}
    num_cols = {"costo_inicial", "dep_acumulada_inicial", "nuevo_costo"}
    text_cols = {"unidad", "unidad_origen", "unidad_destino", "codigo", "codigo_adicional", "cod_rube",
                 "nro_convenio", "ci", "organismo_financiador", "disposicion_respaldo"}

    for spec in SHEETS:
        ws = wb.create_sheet(spec.name)
        ws.sheet_properties.tabColor = "C55A11" if spec.name == "ACTIVOS" else "2E75B6"
        for j, col in enumerate(spec.cols, start=1):
            c = ws.cell(1, j, col.name)
            c.font = HEADER_FONT
            c.fill = REQ_FILL if col.required else HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            c.comment = Comment(("OBLIGATORIO. " if col.required else "Opcional. ") + col.help, "Importador VSIAF")
            letter = get_column_letter(j)
            ws.column_dimensions[letter].width = max(col.width, len(col.name) + 4)
            ref = dv_map.get((spec.name, col.name))
            if ref:
                dv = DataValidation(type="list", formula1="=" + ref, allow_blank=True, showErrorMessage=True,
                                    errorStyle="warning", errorTitle="Valor no listado",
                                    error="El valor no está en el catálogo; el importador lo validará.")
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}5000")
            fmt = "dd/mm/yyyy" if col.name in date_cols else ("#,##0.00" if col.name in num_cols else
                                                              ("@" if col.name in text_cols else None))
            if fmt:
                for r in range(2, 502):
                    ws.cell(r, j).number_format = fmt
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(spec.cols))}1"
        if ejemplos and spec.name in ejemplos:
            for i, rowd in enumerate(ejemplos[spec.name], start=2):
                for j, col in enumerate(spec.cols, start=1):
                    v = rowd.get(col.name)
                    if v is not None:
                        ws.cell(i, j, v)
    # orden final: LEEME, hojas de datos, CATALOGOS
    wb._sheets = [wb["LEEME"]] + [wb[s.name] for s in SHEETS] + [wb["CATALOGOS"]]
    wb.save(path)
