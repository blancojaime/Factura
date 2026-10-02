#!/usr/bin/env python3
"""Extrae las tablas (ListObjects) de los tres libros XLSM originales a JSON crudo.

Uso: python3 extract_seed.py ALMACENES.xlsm COMBUSTIBLE.xlsm ACTIVOS.xlsm [directorio_salida]
Genera almacenes.json, combustible.json y activos.json: { "NombreTabla": [ {col: valor}, ... ] }
Requiere: pip install openpyxl
"""
import json, sys, warnings, datetime as dt
import openpyxl

warnings.filterwarnings("ignore")


def conv(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v if v != "" else None
    return v


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        for t in ws.tables.values():
            rows = list(ws[t.ref])
            head = [str(c.value).strip() for c in rows[0]]
            data = []
            for r in rows[1:]:
                rec = {h: conv(c.value) for h, c in zip(head, r)}
                if any(v is not None for v in rec.values()):
                    data.append(rec)
            out[t.name] = data
    return out


if __name__ == "__main__":
    a, c, f = sys.argv[1:4]
    outdir = sys.argv[4] if len(sys.argv) > 4 else "."
    for name, p in (("almacenes", a), ("combustible", c), ("activos", f)):
        d = extract(p)
        with open(f"{outdir}/{name}.json", "w", encoding="utf-8") as fh:
            json.dump(d, fh, ensure_ascii=False, indent=0, default=str)
        print(name, {k: len(v) for k, v in d.items()})
